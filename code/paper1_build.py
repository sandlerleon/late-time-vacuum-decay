# -*- coding: utf-8 -*-
"""Paper 1 (EPJ C): matter-dilution-triggered recollapse and conditional nonsingular cycles in a metastable dark-energy model.

The homogeneous mechanism only: the matter-held vacuum, its release and the turnaround, the contraction and the effective bounce, reheating, dark
matter and repeated cycles, numerical validation and viability, observational constraints. Localized tunneling and vacuum interfaces are the subject
of the companion paper (Physics of the Dark Universe) and are mentioned only as a distinct extension.

    python paper1_build.py        -> manuscript/EPJC_Dilution_Triggered_Cycles.{docx,tex}, supplement, cover letter (after verify.py)
"""
import datetime
import json
import os
import re
import sys

import numpy as np
from docx import Document
from docx.shared import Pt, Inches

import copyedit as CE
import cyc_blocks as CYB
import dil_blocks as DIL
import render as RN

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, ".."))
OUT = os.path.join(REPO, "manuscript")
FIGDIR = os.path.join(REPO, "figures")
J = lambda f: json.load(open(os.path.join(HERE, f), encoding="utf-8"))
V = J("verification.json")
P1_CHECKS = ["V27_matter_held_vacuum", "V28_hot_cycles", "V29_continuous_cycles", "V30_causal_test", "V31_stability_bounce_relics", "V32_cycle_fixed_point",
             "V33_vacuum_phase"]
P1_CHECKS = [k for k in V if any(k.startswith(p.split("_")[0] + "_") for p in P1_CHECKS)]
if not all(V[k]["passed"] for k in P1_CHECKS):
    sys.exit("verification checks of this paper have not all passed; run verify.py")
DOIS = J("dois.json")
DOI_PAPER, DOI_CODE = DOIS["paper"], DOIS["code"]
COMPANION_DOI = J("companion.json")["paper"] if os.path.exists(os.path.join(HERE, "companion.json")) else None
GITHUB = "https://github.com/sandlerleon/late-time-vacuum-decay"
PM, VM = J("vacuum_phase_map.json"), J("viability_mc.json")
COH = PM["symmetron_coherence"]
REFS = dict(J("refs_all.json"))
TITLE = "Matter-Dilution-Triggered Recollapse and Conditional Nonsingular Cycles in a Metastable Dark-Energy Model"
texsci, sci = CYB.texsci, CYB.sci
GYR = DIL.GYR
RC, RB = DIL.RC, DIL.RB
TV, TV100, TM, TH = DIL.TV, DIL.TV100, DIL.TM, DIL.TH
CR, CAU, CS, BV, RL, PR, CM, NA = CYB.CR, CYB.CAU, CYB.CS, CYB.BV, CYB.RL, CYB.PR, CYB.CM, CYB.NA
SUM, NOEK, EK, A44, COLD, MB = CYB.SUM, CYB.NOEK, CYB.EK, CYB.A44, CYB.COLD, CYB.MB
HELD, YK, MM, LNG, PHB = CYB.HELD, CYB.YK, CYB.MM, CYB.LNG, CYB.PHB
pm, fo, N_CYC = CYB.pm, CYB.fo, CYB.N_CYC
HM = VM["held"]
ISV, RS = J("independent_solver.json"), J("relic_sensitivity.json")

B = []
add = B.append
# ============================================================================================ front matter
ABSTRACT = (
    "Can the dilution of matter end the cosmological expansion and start a new hot phase? We study a scalar field coupled to dark matter through "
    "$A(\\varphi)=1+\\varphi^2/2M^2$, which holds a vacuum with the observed dark-energy density while the dark-matter density exceeds a critical value and "
    "releases it into a negative vacuum once the expansion has diluted the matter. While the field is held, the homogeneous background reproduces "
    "$\\Lambda$CDM to numerical accuracy ($w=-1$); after the release the expansion turns around tens of Gyr later. An intervention that freezes the density seen by the field keeps all %d coupled models of a "
    "parameter grid expanding forever, whereas the uncoupled field with the same potential rolls before today for $\\mu\\ge%.0fH_0$. With an effective "
    "loop-quantum-cosmology bounce, dissipation $\\Upsilon\\propto T^3$ and dark matter that obeys its own Boltzmann equation with a field-dependent mass, %d "
    "consecutive cycles are integrated without resetting any variable. Each has a radiation-dominated bounce, freeze-out to the observed abundance, a matter "
    "era, a held dark-energy epoch of $%s$ e-folds, a release and a turnaround, and the repeating cycle is an attracting fixed point of the bounce-to-bounce "
    "map. The cycle requires a bounce that rethermalizes the dark matter but stays below the percolation of the previous cycle's black holes, dissipation, "
    "inherited rather than regenerated shear and a coherent seed for the release; over stated parameter ranges these hold together in %.1f%% of draws. The "
    "cycle does not produce the observed primordial perturbations ($n_s\\approx3$), so the results are conditional background solutions, not an "
    "observationally complete cosmology."
    % (SUM["runs"], DIL.CTRL_BEFORE[0], N_CYC, pm(HELD), 100 * HM["joint"]))
add(("title", TITLE))
add(("author", "Leon Sandler\nIndependent Researcher, Northbrook, Illinois, USA\nsandler.leon@gmail.com   ·   ORCID 0009-0007-4584-808X"))
add(("abstract", ABSTRACT))
add(("keywords", "cyclic cosmology; nonsingular bounce; dark energy; coupled dark matter; symmetron; recollapse; Boltzmann equation; dissipation; anisotropy"))

# ============================================================================================ 1 introduction
add(("h1", "1. Introduction"))
add(("p", "The idea examined here reverses a common intuition. The dilution of matter is usually the end state of cosmic evolution, the approach to heat "
          "death in an accelerating universe. Here it is the trigger. If the dark energy is the energy of a vacuum held in place by dark matter, the emptying "
          "of the universe destabilizes it, and the expansion can turn into contraction, a bounce and a new hot phase. The paper asks what such a cycle "
          "requires, computes it without prescribed steps, and identifies the conditions under which it fails."))
add(("p", "Cyclic and bouncing cosmologies have a long history \\cite{Tolman,Penrose,GV,ST,IS,NB08,BP15}. Most specify the mechanism that ends the "
          "expansion separately from the observed universe, and many reset some variable between cycles. Three ingredients used here are established. Scalar "
          "fields coupled to matter whose effective potential depends on the local density are known from chameleon and symmetron screening "
          "\\cite{KW,HK,HKLM11} and from coupled dark energy \\cite{Amendola}. An effective bounce at high density follows from loop quantum cosmology "
          "\\cite{APS06,AS11,Singh09}. Thermal dissipation of scalar energy \\cite{Berera95} and thermal freeze-out \\cite{KT90,GG91} are standard early-universe "
          "physics. What is examined here is their combination into a cycle that contains a hot big bang and an epoch like ours, followed without resets, "
          "with control experiments that remove each ingredient and an explicit map of the conditions under which the cycle fails."))
add(("p", "Four results are reported. (i) A symmetron-type coupling to dark matter holds the observed dark energy while the matter density is above a "
          "critical value, with a $\\Lambda$CDM past, and releases it by dilution into a negative vacuum, which ends the expansion; an exact energy relation "
          "fixes how much the dark-matter mass must grow at the release (Sections 2, 3). (ii) Intervention tests show that, within the model, dilution "
          "decides when the expansion ends; a negative vacuum without the coupling also recollapses but cannot reproduce the observed past over most of the "
          "parameter range, so the coupling is not claimed to be the only possible cause of a recollapse (Section 3.3). (iii) With an effective bounce, "
          "thermal dissipation and Boltzmann dark matter, %d consecutive cycles are integrated without any reset; the repeating cycle is an attracting fixed "
          "point of a bounce-to-bounce map, and the length of its dark-energy epoch has a closed form (Section 5). (iv) The conditions are explicit: a bounce "
          "hot enough to rethermalize the dark matter and cold enough to keep black holes from percolating, dissipation, inherited rather than regenerated "
          "anisotropy, and a coherent seed for the release; their joint viability over stated parameter ranges is quantified (Sections 4, 6)." % N_CYC))
add(("h2", "1.1 Relation to previous work"))
add(("p", "Tolman showed that a closed universe that bounces with entropy production grows from cycle to cycle \\cite{Tolman}, and the entropy problem "
          "has remained the central obstacle to cyclic models since: a cycle that returns to the same state must dispose of the entropy it produces, and "
          "geodesically complete cyclic cosmologies face constraints from it \\cite{PS}. The ekpyrotic and cyclic models of Steinhardt and Turok end each "
          "expansion with a scalar that rolls to negative energy and handle the entropy by an exponential expansion that dilutes it \\cite{ST,KOST}; "
          "Penrose's conformal cyclic cosmology matches the remote future of one aeon to the big bang of the next \\cite{Penrose}; pre-big-bang and bouncing "
          "scenarios replace the singularity by a bounce of various origins \\cite{GV,NB08,BP15,IS}. Oscillating models with a bounce show a hysteresis in "
          "which successive cycles grow \\cite{ST12}. The present model shares the negative-vacuum turnaround with the cyclic models and Tolman's growth law, "
          "but differs in three respects. The turnaround is not imposed by a potential that the field rolls down at a prescribed time; it is triggered by the "
          "dilution of the dark matter that held the vacuum, so that the observed dark energy is the held vacuum itself. The entropy is not diluted away but "
          "absorbed in the growth of the next cycle while the local thermal history, set by freeze-out, repeats. And no variable is reset between cycles: the "
          "dark-matter abundance and the restoration of the vacuum are outputs of the equations."))
add(("p", "Couplings of a scalar to matter that make its effective potential density dependent are the basis of chameleon \\cite{KW} and symmetron "
          "\\cite{HK,HKLM11} screening and of coupled dark energy \\cite{Amendola}; in the symmetron the symmetric point is held at high density and broken "
          "at low density, exactly the mechanism used here, but applied to local screening rather than to the cosmological vacuum energy. The effective "
          "bounce of loop quantum cosmology \\cite{APS06,AS11,Singh09} provides a nonsingular continuation at a critical density; it is used here as a "
          "prescription. Thermal dissipation of scalar energy is the mechanism of warm inflation \\cite{Berera95}, and thermal freeze-out \\cite{KT90,GG91} sets "
          "the dark-matter abundance. None of these ingredients is new; their combination into a cycle that contains an epoch like ours is."))
add(("p", "The scope is stated at the outset. The bounce is an effective prescription, not a derivation; the dynamics are homogeneous; and the cycle does "
          "not produce the observed primordial perturbations (Section 7). The results are therefore conditional solutions for the background cosmology. A "
          "distinct possibility, in which dilution makes the held vacuum metastable and it decays by localized bubble nucleation, raises separate questions, "
          "including whether a smooth dark-matter density describes the field on the much shorter length scale of a vacuum wall; it is examined in a "
          "companion paper%s and is not used here." % (" \\cite{Companion}" if COMPANION_DOI else "")))

# ============================================================================================ 2 model
add(("h1", "2. The matter-coupled metastable vacuum"))
add(("h2", "2.1 Action and effective potential"))
add(("p", "The scalar is coupled conformally to dark matter only; baryons and radiation are uncoupled, which avoids the stringent visible-sector "
          "fifth-force tests:"))
add(("eq", r"S=\int d^4x\sqrt{-g}\left[\frac{M_{\rm Pl}^2R}{2}-\frac{1}{2}(\nabla\varphi)^2-V_0(\varphi)\right]+S_{\rm dm}\left[A^2(\varphi)g_{\mu\nu},\psi_{\rm dm}\right]"))
add(("p", "For pressureless dark matter the conserved Einstein-frame density $\\rho_*=\\rho_{*,0}a^{-3}$ enters the field equation through the effective "
          "potential, and the physical dark-matter density is $A(\\varphi)\\rho_*$ \\cite{Amendola,HK}. With"))
add(("eq", r"V_0(\varphi)=V_c-\frac{\mu^2}{2}\varphi^2+\frac{\lambda}{4}\varphi^4,\qquad A(\varphi)=1+\frac{\varphi^2}{2M^2},\qquad V_{\rm eff}=V_0+\rho_*A(\varphi)"))
add(("p", "the effective mass at the symmetric point is $m_{\\rm eff}^2=\\rho_*/M^2-\\mu^2$. The point $\\varphi=0$, with vacuum energy $V_c>0$, is held "
          "while $\\rho_*>\\rho_{\\rm crit}=\\mu^2M^2$ and released when the expansion dilutes the matter below it, at"))
add(("eq", r"a_c=\left(\frac{\rho_{*,0}}{\mu^2M^2}\right)^{1/3},\qquad V_c-V_0(v)=\frac{\mu^2v^2}{4}=\frac{A(v)-1}{2}\,\rho_{\rm crit}"))
add(("p", "This is the symmetron form \\cite{HK,HKLM11} with two changes: the coupling is to dark matter, and the broken vacua $\\varphi=\\pm v$, "
          "$v^2=\\mu^2/\\lambda$, carry negative energy $V_0(v)=-qV_c$. Today $V_c=3\\Omega_\\Lambda H_0^2M_{\\rm Pl}^2$, so the model reproduces the observed "
          "dark energy while it is held, and $a_c>1$ is required for it to be held now. The second relation is exact and has a consequence that cannot be "
          "avoided: the vacuum energy given up at the release equals half the coupling energy stored at the critical density, and since "
          "$\\rho_{\\rm crit}<\\rho_{\\rm dm,0}<V_c$, a release into a negative vacuum needs $A(v)-1>2V_c/\\rho_{\\rm crit}$, i.e. the dark-matter mass must grow "
          "severalfold once the field rolls ($A(v)=%.0f$ for $a_c=2$, $q=1$; $%.1f$--$%s$ over the parameter grid below). At $\\varphi=0$ the coupling "
          "$\\beta=M_{\\rm Pl}\\,d\\ln A/d\\varphi$ vanishes, so the homogeneous background carries no fifth force today. The background equations (flat, "
          "general relativity) are"
          % (DIL.REF.A(DIL.REF.v), min(DIL.AV), DIL.sci(max(DIL.AV), 1))))
add(("eq", r"3M_{\rm Pl}^2H^2=\rho_r+\rho_b+A\rho_*+\frac{\dot\varphi^2}{2}+V_0,\qquad \ddot\varphi+3H\dot\varphi+V_0'+\rho_*A'=0,\qquad \dot\rho_*=-3H\rho_*"))
add(("p", "with $\\dot H=-(\\tfrac43\\rho_r+\\rho_b+A\\rho_*+\\dot\\varphi^2)/2M_{\\rm Pl}^2$. Total stress-energy is conserved: the scalar and the coupled "
          "matter exchange energy through $A'$, and the coupling energy is counted once, in $A\\rho_*$. The equations are integrated from recombination "
          "($z=1100$) with the field at the de Sitter quantum floor, $\\varphi_0=10^{-60}v$ (about $H/2\\pi$ in Planck units), in units of $H_0$ and $M_{\\rm Pl}$ "
          "(dilution_background.py)."))
add(("p", "The comparison with $\\Lambda$CDM has a definite scope, and three levels are distinguished. (i) The homogeneous background: while the field is "
          "held, the background reproduces $\\Lambda$CDM to numerical accuracy (Section 3.1). (ii) Linear perturbations: around $\\bar\\varphi=0$ the field "
          "fluctuation has the mass $m_{\\rm eff}$ and is not sourced by dark-matter density perturbations at first order, because the source term "
          "$\\rho_*A'(\\bar\\varphi)\\delta$ vanishes; the dark matter therefore feels no linear fifth force. This has not been followed through structure "
          "formation, where second-order terms and the time dependence of $m_{\\rm eff}$ enter. (iii) Low-density regions: where the local density stays below "
          "$\\rho_{\\rm crit}$ over more than a Compton length, as in the deepest voids, the held point is locally unstable and the field can roll there before "
          "it does in the mean (Section 7.1). The statement established here is therefore: the homogeneous background reproduces $\\Lambda$CDM to numerical "
          "accuracy while the scalar remains at the symmetric point; perturbative and environmental deviations remain to be tested. The smooth density used "
          "as the source is accurate in this model, because the field's Compton length is cosmological (%.0f Mpc today for $a_c=2$, $\\mu=100H_0$, and of order "
          "a kiloparsec at recombination), so it averages over an enormous number of dark-matter particles."
          % (4431.0 / (100.0 * np.sqrt(7.0)))))
add(("h2", "2.2 Coherence of the release"))
add(("p", "The release is a classical (spinodal) roll: the held point loses stability at $\\rho_{\\rm crit}$ exactly when the broken minima appear, and the "
          "occupied minimum is never metastable (vacuum_phase_map.py). The sign of the roll must therefore be set coherently over the region; otherwise "
          "domain walls between $+v$ and $-v$ form, and they would carry up to %.0f%% of the critical density per Hubble volume. In the cycles of Section 5 the "
          "field restored by the hot phase starts each held epoch %.0f to %.0f times above the Hubble-scale quantum fluctuation $H_\\Lambda/2\\pi$, a marginal "
          "margin; any explicit breaking of $\\varphi\\to-\\varphi$ with a tilt $\\epsilon\\gtrsim%s\\,\\mu^2v$ sets the sign coherently, and this is assumed. "
          "Matter inhomogeneity does not produce sharp interfaces: the field averages the density over its Compton length (%.0f Mpc at release), where linear "
          "fluctuations are $\\sigma_R=%.3f$, so the release time varies by $\\delta a/a\\approx%.3f$ between such regions and the homogeneous treatment holds "
          "to that accuracy."
          % (100 * COH["domain_walls"][-1]["wall_energy_per_Hubble_volume_over_rho_crit"], COH["phi_rest_over_v"]["min"] / COH["dphi_quantum_over_v"],
             COH["phi_rest_over_v"]["max"] / COH["dphi_quantum_over_v"], texsci(COH["eps_hat_min"], 0), COH["compton_scale_Mpc"], COH["sigma_R_linear_today"],
             COH["release_spread_delta_a_over_a"])))

# ============================================================================================ 3 transition and turnaround
add(("h1", "3. Density-triggered transition and cosmological turnaround"))
add(("h2", "3.1 Release and recollapse"))
add(("p", "Three models are compared with identical initial data: $\\Lambda$CDM (A), the uncoupled scalar with the same potential (B, the control) and the "
          "coupled scalar (C). For the reference case $a_c=2$, $\\mu=100H_0$, $q=1$ (Fig. 1), $\\Lambda$CDM expands forever. The coupled field stays at "
          "$|\\varphi|/v=%s$ up to the present, so the homogeneous past history is that of $\\Lambda$CDM ($w_0=-1$ to all digits; $H(a)$ agrees with $\\Lambda$CDM to the "
          "integration accuracy), is released after $a_c$, rolls at $a=%.2f$ and turns the expansion around at $a=%.2f$, %.0f Gyr from now. The uncoupled "
          "control has been unstable since recombination and turns around at $a=%.2f$, %.0f Gyr from now."
          % (DIL.sci(RC["phi_today_over_v"], 1), RC["a_roll"], RC["a_turn"], RC["t_turn"] * GYR, RB["a_turn"], RB["t_turn"] * GYR)))
add(("p", "Over a grid of %d coupled models ($a_c=1.25$--$5$, $\\mu=3$--$300H_0$, $q$ from $-0.5$ to $10$) every one reaches the present with the vacuum "
          "held; %d release into a negative vacuum and recollapse, and the %d with $q<0$ settle in a lower de Sitter vacuum and expand forever (Fig. 2b, "
          "Table 1). The control rolls before today for $\\mu\\ge%.0fH_0$, so it cannot describe the observed universe there, while the coupled model can at every "
          "$\\mu$. The coupling is what makes a late release compatible with an observed $\\Lambda$CDM past. The release is not instantaneous: the field grows "
          "from the quantum floor at a rate of order $\\mu$, so the roll lags $a_c$ by about $\\ln(v/\\varphi_0)/\\mu$ Hubble times, from $a=%.2f$ to $%.2f$ for "
          "seeds $10^{-6}$ to $10^{-60}v$; the control rolls before today for every seed above $10^{-40}v$ (Fig. 2c)."
          % (DIL.N_C, DIL.N_REC, DIL.N_NOREC, DIL.CTRL_BEFORE[0], DIL.SEC[1e-6]["a_roll"], DIL.SEC[1e-60]["a_roll"])))
add(("p", "The release changes the matter content as well as the vacuum. Once the field rolls to $v$ the dark-matter mass grows by $A(v)$, so the "
          "physical dark-matter density jumps by that factor while the vacuum energy falls from $V_c$ to $-qV_c$; the energy relation of Section 2.1 makes the "
          "two changes balance at the release. The universe then decelerates, and the turnaround occurs when the diluting matter and radiation no longer "
          "compensate the negative vacuum, $A(v)\\rho_*+\\rho_b+\\rho_r\\simeq qV_c$, within a fraction of a Hubble time of the release for $q=O(1)$; for "
          "$q<0$ the field settles in a lower de Sitter vacuum and the expansion continues. The model therefore makes the timing of the end of the "
          "expansion a property of the matter density, not of a clock in the potential: the release occurs when $\\rho_*$ falls to $\\rho_{\\rm crit}$, "
          "whatever the history before it."))
rowsT = []
for mu in sorted({r["mu"] for r in DIL.ROWS}):
    rc_, rb_ = DIL.row(2.0, mu, 1.0, True), DIL.row(2.0, mu, 1.0, False)
    rowsT.append(["%g" % mu, "%.3g" % rc_["a_roll"], "%.3g" % rc_["a_turn"], "%.0f" % (rc_["t_turn"] * GYR),
                  ("rolled before today" if rb_["cls"] == "rolled before today" else "%.3g" % rb_["a_turn"]), "%.0f" % rc_["A_at_v"]])
add(("table", "Table 1. Release and turnaround of the coupled model and the uncoupled control ($a_c=2$, $q=1$, seed $10^{-60}v$ at $z=1100$). Times from "
              "today; $A(v)$ is the factor by which the dark-matter mass grows after the release.",
     ["$\\mu/H_0$", "roll, $a$", "turnaround, $a$", "turnaround (Gyr)", "control turnaround, $a$", "$A(v)$"], rowsT, None))
add(("fig", "fig13_matter_held_vacuum.png", "Fig. 1. The matter-held vacuum ($a_c=2$, $\\mu=100H_0$, $q=1$). (a) Effective potential for several matter "
            "densities; the symmetric point, with energy $V_c$, is a minimum while $\\rho_*>\\rho_{\\rm crit}$. (b) Hubble rate against time from today for the "
            "coupled model, the uncoupled control and $\\Lambda$CDM. (c) Field amplitude from recombination: the coupled field is held at the quantum floor until "
            "$a_c$, the control grows from the start.", 6.3))
add(("h2", "3.2 Cooling or dilution?"))
add(("p", "A thermal mass $cT^2\\varphi^2/2$ from a bath of temperature $T$ acts in the same way: the symmetric point is held while "
          "$\\rho_*/\\rho_{\\rm crit}+(T/T_c)^2>1$, $T_c=\\mu/\\sqrt c$, and the possible triggers are compared in Fig. 2a. With the critical scale factor "
          "set to 2 for each, matter alone releases the field at $a=2$, temperature alone at $a=2$, both together at $a=%.2f$, and all three recollapse "
          "($a=%.1f$, $%.1f$ and $%.1f$). Mathematically, then, cooling could trigger the release; physically it cannot. In physical units $\\sqrt c=\\mu a_T/T_0$, "
          "so a release at $a_T\\sim2$ with $\\mu=100H_0$ needs $\\sqrt c=%s$, and a scalar with so weak a coupling interacts with the bath at a rate "
          "$\\sim c^2T\\approx%s\\,H_0$: it never reaches equilibrium, and the thermal mass is never established. A coupling strong enough for equilibrium "
          "today, $\\sqrt c>%s$, holds the field until $a_T\\sim%s$. The two requirements exclude each other by twenty orders of magnitude. Cooling does not "
          "trigger the release; dilution of dark matter can."
          % (TH["a_release_C_analytic"], TM["A_matter"]["a_turn"], TM["B_temperature"]["a_turn"], TM["C_both"]["a_turn"], DIL.sci(TV100["sqrt_c_needed"], 1),
             DIL.sci(TV100["gamma_over_H0"], 0), DIL.sci(TV["sqrt_c_for_equilibrium_today"], 0), DIL.sci(TV["a_T_if_equilibrated"]["100.0"], 0))))
add(("fig", "fig14_release_scan.png", "Fig. 2. (a) The plane of matter density and temperature: the symmetric point is held above the line; trajectories of "
            "the matter-only, temperature-only and combined triggers (dots: release). (b) Scale factor at turnaround against $\\mu$ for the coupled model at "
            "five $a_c$ and for the uncoupled control, which rolls before today at $\\mu=300H_0$. (c) Scale factor of the release against the initial seed for "
            "the coupled model and the control.", 6.3))
add(("h2", "3.3 Interventions on the turnaround"))
add(("p", "Section 3.1 compared the coupled model with an uncoupled field. A sharper test intervenes on the cause itself (causal_controls.py). Five models "
          "are evolved from today over the grid of Table 1: P, a positive cosmological constant (expands forever); N, a constant negative vacuum with matter "
          "(recollapses whatever the matter does, but decelerates today, $q_0=%.2f$ for $q=1$, and is excluded by the observed acceleration); U, the scalar "
          "without the coupling; F, the coupled model with the matter density seen by the field frozen at its present value (an intervention on the field "
          "equation only, $do(\\rho_*=\\rho_{*,0})$; gravity still sees diluting matter); and D, the model of this paper. In all %d grid points F stays held "
          "and expands forever, since $m_{\\rm eff}^2=\\mu^2(a_c^3-1)>0$, while D recollapses for every $q>0$ (%d of %d) and releases but keeps expanding for "
          "$q<0$ (there $V_0\\ge-qV_c>0$, so $\\rho>0$ and $H$ cannot vanish); U is never held: it rolls, or for $\\mu=3H_0$, $q=10$ recollapses before "
          "reaching $v/2$ (Fig. 3). Within the model, the outcome changes from eternal expansion to recollapse when, and only because, the matter seen by the "
          "field is allowed to dilute. This does not make the coupling the only possible cause of a recollapse: a negative vacuum without it (N, U) also "
          "recollapses, but with a past that is not that of the observed universe over most of the range."
          % (CAU["N"]["1.0"]["q0_deceleration"], SUM["runs"], SUM["dilution_causes_recollapse"], SUM["q_positive"])))
add(("fig", "fig19_causal_controls.png", "Fig. 3. Causal test. (a) $H(t)$ for the five models at $a_c=2$, $\\mu=30H_0$, $q=1$: only D turns around after a "
            "$\\Lambda$CDM past; F, identical except that the matter seen by the field does not dilute, expands forever. (b) Outcome over the grid of $a_c$, "
            "$\\mu$ and $q$.", 6.3))

# ============================================================================================ 4 contraction and bounce
add(("h1", "4. Contracting evolution and the effective bounce"))
add(("h2", "4.1 The bounce prescription and why dissipation is needed"))
add(("p", "A recollapse becomes a cycle only if the contraction reverses, which needs physics beyond general relativity. The effective "
          "loop-quantum-cosmology equations for a flat universe, for which they are derived \\cite{APS06,AS11},"))
add(("eq", r"H^2=\frac{\rho}{3M_{\rm Pl}^2}\left(1-\frac{\rho}{\rho_c}\right),\qquad \dot H=-\frac{\rho+p}{2M_{\rm Pl}^2}\left(1-\frac{2\rho}{\rho_c}\right)"))
c4, c5, c6 = DIL.cold(1e4), DIL.cold(1e5), DIL.cold(1e6)
gmed = lambda r: DIL.gm([c.get("G") for c in r["cycles"][1:]])
add(("p", "are used as the bounce, with no curvature term appended. This is an assumption of the paper: the bounce is prescribed, not derived, and the "
          "physical $\\rho_c\\sim\\rho_{\\rm Pl}$ cannot be integrated, so $\\rho_c$ is set far below it and the dependence on it is reported. A bounce alone is "
          "not enough. Without dissipation the contraction is dominated by the scalar's kinetic energy, which grows as $a^{-6}$, faster than radiation, and the "
          "bounce occurs at low density and is cold (cyclic_extension.py): at $\\rho_c=10^4V_c$ the cycles repeat with $G_n\\simeq%.2f$ but have no hot dense "
          "epoch, and they grow ($G_n\\simeq%.2f$ and $%.2f$) at $\\rho_c=10^5$ and $10^6V_c$, the cosmological hysteresis known from oscillating models "
          "\\cite{ST12}. Thermal dissipation, $\\Upsilon=C_TT^3$, the form used for dissipative scalar dynamics in warm inflation \\cite{Berera95}, is negligible "
          "in the cold expanding universe and dominant in the hot contraction, where it turns the scalar's energy into radiation, so that the bounce is "
          "radiation dominated and hot. Sections 4.2--4.5 test the stability of repeated cycles, the contraction and the bounce; Section 5 follows the full "
          "cycle."
          % (gmed(c4), gmed(c5), gmed(c6))))
import osc_blocks as OSC
add(("h2", "4.2 Stability of repeated cycles"))
add(("p", "Why dissipation is required can be isolated in the simplest homogeneous cycle: an open universe in a negative vacuum, $V=V_t+m^2\\varphi^2/2$ with "
          "$V_t<0$, with radiation and the effective bounce (oscillation.py, oscillation_validation.py). The hierarchy $\\rho_c/|V_t|$ cannot be integrated at its "
          "physical value, so the cycles are simulated at $\\rho_c/|V_t|=10^4$--$10^8$ in Planck units and the scaling of each quantity is measured. With radiation "
          "alone the cycles are exactly periodic: over %d bounces the period varies by %s, it scales as $|V_t|^{%.3f}$ and does not depend on $\\rho_c$, and "
          "the trajectory is far from a sine (the fundamental harmonic carries %.1f%% of the power of $H$, and $\\dot H>0$ only during %s of each cycle, at the "
          "bounce). Any oscillation of the scalar, however, is amplified from cycle to cycle: in the linear regime its energy grows by a factor between %.0f and "
          "%.0f per cycle, because the field is frozen near the bounce and released with a phase that does not return the energy it gained. Once it is no longer "
          "small, successive cycles grow, the cosmological hysteresis of oscillating models \\cite{ST12}, and the growth obeys an energy-balance law: the "
          "continuity equation gives $\\Delta(\\rho a^3)=-W_n$ with $W_n=\\oint p\\,d(a^3)$, and at a turnaround $\\rho a^3=-3a_{\\max}$, so that "
          "$\\Delta a_{\\max,n}=W_n/3$, satisfied to %s."
          % (OSC.PER["n_bounces"], OSC.sci(OSC.PER["period_spread"], 0), OSC.O["scaling"]["period_slope_in_Vt"], 100 * OSC.PER["fundamental_power_fraction"],
             OSC.sci(OSC.PER["fraction_of_cycle_with_Hdot_positive"], 1), min(OSC.MU), max(OSC.MU), OSC.sci(max(OSC.LAWS.values()) + 1e-16, 0))))
add(("p", "Dissipation into radiation competes with the amplification. The cycles run away for $\\Upsilon P\\le%.2f$ and are stabilized for "
          "$\\Upsilon P\\ge%.1f$, consistent with the criterion $\\Upsilon P\\gtrsim\\ln\\mu=%.2f$ ($P$ the period, $\\mu$ the energy multiplier). Over a grid of "
          "masses, dissipation strengths and initial amplitudes, %d of %d solutions are stable after 100 or more cycles and the rest grow; none shrinks and none "
          "loses its turnaround, and the stable ones remain stable over 1000 cycles, with $|G_n-1|\\le%s$ over the last 500 (Fig. 9). Stabilized cycles keep a "
          "constant entropy after a one-time increase; runaway cycles produce entropy without bound, which is the entropy problem of cyclic models "
          "\\cite{Tolman,PS}. The instability does not depend on the bounce prescription (with $H^2=(\\rho/3)[1-(\\rho/\\rho_c)^2]$ the multiplier is %.1f). Every "
          "cycle of this simple model has negative vacuum energy and therefore no epoch like ours; the matter-held vacuum of Section 2 is what supplies one, and "
          "the dissipation $\\Upsilon=C_TT^3$ of Section 5 is what keeps its cycles on the stable side of this boundary."
          % (max(OSC.RUN), min(OSC.STAB), OSC.LNMU, sum(v for k, v in OSC.NCLS.items() if k == "stable"), len(OSC.PD),
             OSC.sci(max(l_["max_abs_G_minus_1_last_500"] for l_ in OSC.LONG), 0), OSC.ROB["n=2"]["multiplier"])))
add(("fig", "fig12_hysteresis_stability.png", "Fig. 9. Stability of repeated cycles in a negative vacuum with an effective bounce. (a) Change of the turnaround "
            "radius against the pressure-volume loop $W_n/3$ for curvature-dominated cycles; dashed: the law $\\Delta a_{\\max}=W/3$. (b) Classification after "
            "100 or more cycles over the scalar mass and the dissipation (circles $\\varphi_0=10^{-8}$, squares $10^{-3}$): stable above $\\Upsilon P\\approx\\ln\\mu$. "
            "(c) A vacuum transition at the third turnaround turns a stable sequence into a growing one.", 6.3))
add(("h2", "4.3 Anisotropy in the contraction"))
SHB = CYB.SHB
add(("p", "A contracting universe is unstable to anisotropy: for Bianchi I the shear enters the Friedmann equation as $\\rho_\\sigma=\\Sigma/a^6$, which grows "
          "relative to radiation as $a^{-2}$. In a single contraction from the turnaround to a bounce at $\\rho_c$ the shear stays below 10%% of the density at the "
          "bounce only if it is below $s_{\\max}=$%s of the matter density at the turnaround (contraction_stability.py; the analytic scaling is reproduced to %s). "
          "The cycles of Section 5 meet this condition in a way specific to them: the entropy produced in each cycle enlarges the next, so at a fixed "
          "temperature the scale factor is larger by $G_n$ and the shear, which is conserved in comoving terms, is diluted by $G_n^6$ per cycle. In the %d-cycle "
          "run the shear fraction at the bounce falls from %s to %s and then below %s. This controls the shear inherited from the past, not shear generated "
          "anew in each cycle; for that the margin is the tolerance $s_{\\max}$ itself."
          % (sci(NOEK["s_max"], 1), "%.1g" % abs(CS["analytic_check"]["numerical"] / CS["analytic_check"]["analytic"] - 1), N_CYC, sci(max(SHB[0], 1e-300), 0),
             sci(max(SHB[1], 1e-300), 0), sci(max(min(SHB[2:]) if len(SHB) > 2 else 1e-300, 1e-300), 0))))
add(("p", "An ekpyrotic phase widens the tolerance. A second scalar $\\chi$ with $V_{\\rm ek}=-V_ee^{-c\\chi/M_{\\rm Pl}}$, $c^2>6$, has $w=c^2/3-1>1$ and "
          "$\\rho_\\chi\\propto a^{-c^2}$, so it dilutes the shear as $a^{c^2-6}$ \\cite{KOST,EWST04}. With its own dynamics, an exit and a decay into radiation "
          "after the exit, and starting with $\\rho_\\chi=10^{-12}\\rho_m$ at the turnaround, $s_{\\max}$ rises from %s to %s ($c=3$), %s ($c=5$) and %s ($c=7$), "
          "with %.1f, %.1f and %.1f e-folds of ekpyrotic contraction (Fig. 4). The tolerance is set by the onset: the ekpyrotic phase must dominate before the "
          "shear does. Ekpyrosis is not part of the cycle, which does not need it for the inherited shear; it is the mechanism the model would need if "
          "anisotropy were regenerated at the level $s>$%s each cycle."
          % (sci(NOEK["s_max"], 1), sci(EK[0]["s_max"], 1), sci(EK[1]["s_max"], 1), sci(EK[2]["s_max"], 1), EK[0]["ekpyrotic_efolds"], EK[1]["ekpyrotic_efolds"],
             EK[2]["ekpyrotic_efolds"], sci(NOEK["s_max"], 0))))
add(("fig", "fig20_contraction_stability.png", "Fig. 4. Shear in the contraction. (a) Shear fraction at the bounce against the anisotropy at the turnaround, "
            "without and with an ekpyrotic phase ($f_0=10^{-12}$, exit at $10^{-4}\\rho_c$). (b) The largest controlled anisotropy against the ekpyrotic "
            "fraction at the turnaround, for three exit densities ($c=5$).", 6.3))
add(("h2", "4.4 Bounce validation"))
LQC, GH = CYB.LQC, CYB.GH
MEV = CYB.MEV
add(("p", "Three tests address the effective bounce (bounce_validation.py). (A) Sensitivity to $\\rho_c$: bounces at %s MeV behave identically up to the "
          "bounce; the next cycle depends on $\\rho_c$ only through rethermalization (next cycle: %s). (B) The effective equations through the bounce: without "
          "projection inside $\\rho>0.4\\rho_c$ the constraint holds to %s, $d\\ln\\rho/d\\ln a=-3(1+w)$ to %s, the density and the entropy are continuous across "
          "the bounce, and the curvature invariants are finite, $R=%.4f\\rho_c$ and $K=%.4f\\rho_c^2$ at the bounce against the analytic $3(\\rho+p)=4\\rho_c$ and "
          "$16\\rho_c^2/3$ for radiation. (C) An independent bounce: general relativity with a negative-energy stiff fluid $\\rho_g=-g/a^6$ in place of the LQC "
          "correction. Its bounce density is set by a comoving constant, so it is not fixed from cycle to cycle, and the bounce temperature is %s MeV in "
          "successive cycles, against %s MeV for the LQC bounce. The cycle survives the change of mechanism, but its self-similarity is a property of a bounce "
          "at a fixed density, not of bounces in general (Fig. 5)."
          % (", ".join("%.3g" % a["T_bounce_MeV"] for a in BV["A"]),
             "; ".join("%.3g MeV: %s" % (a["T_bounce_MeV"], "working" if (fo(a["next"]) or 0) > 0.1 and (a["next"].get("held_accelerated_efolds") or 0) > 0.5 else "empty") for a in BV["A"]),
             sci(max(a["max_residual"] for a in BV["A"]), 0), sci(max(a["conservation_median_error"] for a in BV["A"] if a["conservation_median_error"]), 0),
             A44["R_over_rho_c"], A44["K_over_rho_c2"], ", ".join("%.3g" % (c["T_bounce"] * MEV) for c in GH["cycles"] if c.get("T_bounce")),
             ", ".join("%.3g" % (c["T_bounce"] * MEV) for c in LQC["cycles"] if c.get("T_bounce")))))
add(("fig", "fig21_bounce_validation.png", "Fig. 5. Bounce validation. (a) Curvature invariants and density through the first bounce of the baseline run; "
            "dashed lines: the analytic bounce values for radiation. (b) Test C: bounce temperature per cycle for the LQC bounce and for a ghost stiff "
            "fluid.", 6.3))
add(("h2", "4.5 Black holes through the bounce"))
perc = RL["percolation"]
P_ = CR["params"]
add(("p", "The black holes of a cycle do not take part in its rethermalization. Their mass function $n(M,t)$ is evolved through the contraction "
          "(compact_relics.py) with $\\partial_tn+3Hn+\\partial_M(\\dot Mn)=C_{\\rm merge}+S_{\\rm form}$: accretion of radiation $\\dot M=27\\pi G^2M^2\\rho_r$ "
          "\\cite{CH74}, capped at the Hubble mass; binned mergers with the capture cross-section; and the present population as the source (stellar remnants "
          "with $\\Omega=10^{-4}$ near $10\\,M_\\odot$, supermassive holes with $\\Omega=7.6\\times10^{-6}$ near $10^8M_\\odot$, with exponential cut-offs above "
          "100 and $5\\times10^9M_\\odot$). The quantity that decides whether a homogeneous bounce is possible is the fraction of space inside horizons, $f_V$. "
          "At the bounce ($T_b=%.0f$ MeV) it is %s for the present population and %s with accretion and mergers (Fig. 6a). Holes above $%.0f\\,M_\\odot$ "
          "accrete runaway to the Hubble mass, and since $f_V\\propto T^3$ the evolved population would reach $f_V=1$ for a bounce above about %.1f GeV; "
          "all supermassive mass in holes of a single mass gives %s MeV for $10^{10}$, $10^9$ and $10^8M_\\odot$. The black holes therefore bound the bounce "
          "from above, while BBN bounds it from below: $5\\ \\mathrm{MeV}\\lesssim T_b\\lesssim%.1f$ GeV within this prescription, or $\\lesssim%.1f$ GeV for "
          "the single-mass estimate. The surviving holes enter the next cycle as relics, and their number per unit entropy falls by $G_n^3$ per cycle."
          % (RL["T_bounce_MeV"], sci(RL["E2"]["f_V_bounce"], 0), sci(RL["E3"]["f_V_bounce"], 0), RL["M_runaway_Msun"], RL["T_perc_E3_MeV"] / 1e3,
             ", ".join("%.3g" % perc[k]["T_MeV"] for k in ("SMBH 1e10 Msun", "SMBH 1e9 Msun", "SMBH 1e8 Msun")), RL["T_perc_E3_MeV"] / 1e3,
             perc["SMBH 1e10 Msun"]["T_MeV"] / 1e3)))
RSB = {r["variant"]: r for r in RS["rows"]}
POPV = [r for r in RS["rows"] if r["variant"] not in ("baseline (E3)",) and "accretion" not in r["variant"] and "E2" not in r["variant"]
        and "mergers" not in r["variant"]]
add(("p", "This upper limit is an estimate within the stated population prescription, not a general cosmological bound, and its sensitivity was tested "
          "(relic_sensitivity.py, Table S9). Adding horizon volumes treats the holes as nonoverlapping; for randomly placed, overlapping spheres the occupied "
          "fraction is $1-e^{-f_V}$ and continuum percolation sets in at $f_V\\approx%.2f$, which lowers the bound from %.2f to %.2f GeV. Changing the stellar "
          "abundance threefold either way leaves it unchanged; the supermassive abundance, cut-off and characteristic mass move it between %.2f and %.2f GeV "
          "(between %.2f and %.2f GeV with the overlap threshold); without accretion and mergers it is %.2f GeV, and without mergers alone it is unchanged. "
          "Accretion is the sensitive input. With two and three times the geometric capture rate the bound falls to %.2f and %.2f GeV, and with ten times to "
          "%.0f MeV, below the baseline bounce, because the mass above which holes grow to the Hubble mass falls with the rate. The geometric capture "
          "cross-section is the physical rate for radiation, and Bondi accretion of a relativistic fluid is of the same order, so the baseline is the "
          "physical choice, but the bound inherits that assumption. Clustering lowers it locally: the supermassive holes trace galaxies, and a region "
          "overdense by $1+\\delta$ reaches the threshold at a temperature lower by $(1+\\delta)^{-1/3}$, %.2f GeV for $\\delta=100$. A strongly "
          "inhomogeneous contraction, in which horizon volumes cannot be treated as a dilute gas at all, is outside this estimate. Within these assumptions "
          "the window is about 5 MeV to %.1f--%.1f GeV, and the baseline bounce at %.0f MeV lies inside it for every variant except the tenfold accretion. "
          "Closing the window at the overlap threshold, at the lowest population variant or at the clustered estimate lowers the joint viability of "
          "Section 6.3 from %.1f%% to %.1f%%, %.1f%% and %.1f%%."
          % (RS["f_c_overlap"], RSB["baseline (E3)"]["T_perc_GeV_fc1"], RSB["baseline (E3)"]["T_perc_GeV_overlap"],
             min(r["T_perc_GeV_fc1"] for r in POPV), max(r["T_perc_GeV_fc1"] for r in POPV), min(r["T_perc_GeV_overlap"] for r in POPV),
             max(r["T_perc_GeV_overlap"] for r in POPV), RSB["no accretion, no mergers (E2)"]["T_perc_GeV_fc1"], RSB["accretion x 2"]["T_perc_GeV_fc1"],
             RSB["accretion x 3"]["T_perc_GeV_fc1"], 1e3 * RSB["accretion x 10"]["T_perc_GeV_fc1"], RS["clustered_overlap_GeV"]["100.0"],
             RS["physical_accretion_range_GeV"][0], RS["physical_accretion_range_GeV"][1], RL["T_bounce_MeV"],
             100 * RS["held_viability"]["baseline estimate"]["joint"], 100 * RS["held_viability"]["overlap threshold"]["joint"],
             100 * RS["held_viability"]["lowest population variant, overlap"]["joint"], 100 * RS["held_viability"]["region overdense by 100, overlap"]["joint"])))
add(("fig", "fig22_relics_failure_map.png", "Fig. 6. (a) Fraction of space inside black-hole horizons through the contraction, for the present population "
            "(E2) and with accretion and mergers (E3). (b) Outcome of five continuous cycles over dark-matter mass and bounce temperature: the cycles work "
            "where the bounce rethermalizes the dark matter (Section 6.2).", 6.3))
add(("h2", "4.6 Physical status of the low-density bounce"))
add(("p", "The bounce used in this paper is a phenomenological prescription, and its status should be stated plainly. Loop quantum cosmology derives "
          "the effective equations of Section 4.1 for a bounce at $\\rho_c\\approx0.41\\rho_{\\rm Pl}$ \\cite{APS06,AS11}. The simulations place it instead at "
          "$\\rho_c=10^{44}H_0^2M_{\\rm Pl}^2$, a bounce temperature of about %.0f MeV, some twenty orders of magnitude below the Planck temperature and some eighty "
          "below the Planck density. At that density standard loop quantum cosmology predicts no bounce, and no established theory predicts one. The quadratic "
          "form $H^2=(\\rho/3)(1-\\rho/\\rho_c)$ is used because it is the simplest modification of the Friedmann equation that is nonsingular, conserves "
          "stress-energy and fixes the bounce at a density; the instability of repeated cycles does not depend on this form (Section 4.2). The "
          "ghost-fluid bounce of Section 4.4 is a sensitivity test, not a physical proposal: a negative-energy stiff fluid violates the null energy condition "
          "and is not known to be a healthy field theory." % RL["T_bounce_MeV"]))
add(("p", "What the simulations establish is therefore conditional. If some mechanism reverses the contraction at a fixed density between the BBN and "
          "black-hole bounds, the cycle of Section 5 follows from the remaining equations, and the properties that depend on the bounce are isolated: the bounce "
          "temperature is set by $\\rho_c$; the next cycle depends on $\\rho_c$ only through the rethermalization of the dark matter (Section 4.4); and "
          "self-similar cycles need a bounce at a fixed density, which a ghost-fluid bounce does not provide. The simulations are not a derivation of a "
          "realistic quantum-gravity bounce, and none of the results should be read as evidence that such a bounce occurs."))

# ============================================================================================ 5 cycles
add(("h1", "5. Reheating, dark matter and repeated cycles"))
add(("h2", "5.1 Continuous reheating without resets"))
add(("p", "The cycle is followed with no variable reset (continuous_cycles.py). The dark matter is a thermal species with its own Boltzmann equation "
          "\\cite{KT90,GG91}; its particles have the field-dependent mass $m_{\\rm eff}=mA(\\varphi)$ and are in kinetic equilibrium with the bath, so that their "
          "comoving number $N=na^3$ and their energy per particle follow"))
add(("eq", r"\dot N=-\langle\sigma v\rangle\frac{N^2-N_{\rm eq}^2}{a^3},\qquad n_{\rm eq}=c_n\frac{x^2K_2(x)}{2}T^3,\qquad \epsilon=m_{\rm eff}\frac{K_1(x)}{K_2(x)}+3T,\qquad x=\frac{m_{\rm eff}}{T}"))
add(("p", "The scalar exchanges energy with the bath through the dissipation and through the field dependence of the particle masses:"))
add(("eq", r"\ddot\varphi+(3H+\Upsilon)\dot\varphi+V_0'+F=0,\qquad F=n\frac{K_1}{K_2}mA'(\varphi),\qquad \dot\rho_{\rm th}=-3H(\rho_{\rm th}+p_{\rm th})+\Upsilon\dot\varphi^2+F\dot\varphi"))
add(("p", "with $\\rho_{\\rm th}=\\rho_r+n\\epsilon$, $p_{\\rm th}=\\rho_r/3+nT$ and $\\Upsilon=C_TT^3$; the temperature is a state variable and the total "
          "stress-energy is conserved by construction. The single coupling $A(\\varphi)$ produces both the hot-phase thermal mass (a relativistic gas gives "
          "$F\\simeq nm_{\\rm eff}mA'/2T$) and the late hold (a cold gas gives $F=\\rho_*A'$), so the restoration of the vacuum and the freeze-out of the dark "
          "matter are outputs. In physical units $T=1$ is 1.82 meV, $C_T=1$ is a dissipation $T^3/f^2$ with $f=(H_0M_{\\rm Pl}^3)^{1/4}=2.1$ TeV, and "
          "$\\langle\\sigma v\\rangle$ is fixed once, by requiring the present dark-matter density in the held epoch of the repeating cycle ($%s$ cm$^3$ "
          "s$^{-1}$ for $m=%s$ MeV). This is %s times its value for freeze-out at constant mass, because in the cycle the dark matter freezes out while the "
          "field is still displaced and the dark matter heavier. Two numerical devices are used and tested in Section 6.1: where the dissipation exceeds every "
          "other rate by $10^3$ the field moves at its terminal velocity, and $H$ is projected onto the constraint once per doubling of the scale factor away "
          "from the turning points; relative to the density the constraint residual stays below %s throughout."
          % (texsci(PR.get("sigv_cm3s", 6.35e-26), 1), MB, "%.3g" % CR["params"].get("sigv_scale", 1.0), sci(CR["max_residual"], 0))))
add(("p", "The first result is negative. With $m=%s$ MeV and a bounce at %.1f MeV, our cycle completes (release, turnaround at $\\ln a=%.2f$, a "
          "radiation-dominated bounce) but the next cycle is empty: the dark matter reaches the bounce with $x=m_{\\rm eff}/T=%.0f$, is annihilated during the hot "
          "contraction and is not regenerated ($Y/\\kappa=%s$ in the next cycle), the field is frozen at $|\\varphi|=%.2fv$ by the overdamping, and there is no "
          "matter era (peak matter fraction %.3f) and no dark-energy epoch. Prescribed resets of the dark-matter abundance at freeze-out, used in an earlier "
          "version of this model, hide exactly this failure."
          % (MB, COLD["T_bounce_MeV"], CR["cycles"][0]["ln_a_max"], COLD["bounce_x"], texsci(max(fo(COLD["next"]), 1e-30), 1), COLD["next"]["phi_bounce_over_v"],
             COLD["next"]["max_matter_fraction"])))
add(("p", "The second result is that the same equations give working cycles once the bounce rethermalizes the dark matter, $T_b\\gtrsim m_{\\rm eff}/5$. "
          "With $m=%s$ MeV, a bounce at %.0f MeV and today's anisotropy set to $\\Omega_\\sigma=10^{-30}$, %d cycles were integrated with no reset of any "
          "variable (Fig. 7, Table 2). In every cycle the field is frozen at the bounce at $|\\varphi|/v=%s$, the dark matter returns to equilibrium "
          "($x_b=%.1f$--$%.1f$) and freezes out while the field is still displaced (yield $Y_{\\rm fo}/\\kappa=%s$); the field is then restored, which lowers "
          "the dark-matter mass and heats the bath, and the held epoch has $Y/\\kappa=%s$; the field is damped toward $\\varphi=0$ by %s orders of magnitude, a "
          "matter era follows (peak matter fraction $%s$), then a dark-energy epoch with the field held ($%s$ e-folds of accelerated expansion), a release, a "
          "turnaround and the next bounce. The held e-folds vary between %.2f and %.2f, because the value at which the field freezes depends on the phase of its "
          "oscillation when the overdamping sets in (Section 5.2). The growth per cycle, $\\ln G_n=%s$, equals the entropy growth "
          "$\\frac{1}{3}\\ln(S_{n+1}/S_n)$ to within %s in every cycle: the second law holds and is absorbed by Tolman's growing cycles \\cite{Tolman}."
          % (MB, A44["T_bounce_MeV"], N_CYC, pm(PHB), min(c.get("bounce_x", 0) for c in CR["cycles"]), max(c.get("bounce_x", 0) for c in CR["cycles"]),
             pm(CYB.arr("Y_fo_over_kappa")), pm(YK), CYB.PHMIN_DEC, pm(MM), pm(HELD), HELD.min(), HELD.max(), pm(LNG), sci(CYB.GDEV, 0))))
if CYB.C100:
    c1 = [c for c in CYB.C100["cycles"] if c["n"] >= 1 and "held_accelerated_efolds" in c]
    a1 = lambda k: np.array([c[k] for c in c1 if c.get(k) is not None], dtype=float)
    add(("p", "A second run of %d cycles with $m=100$ MeV and the same bounce (continuous_cycles_100MeV.json) works as well: held e-folds $%s$, peak matter "
              "fraction $%s$, $\\ln G_n=%s$, with the field freezing at $\\varphi_b=%sv$. That point is marginal: it lies on the working branch of Section 6.2 "
              "only because the field freezes far from $v$, and one-at-a-time changes of $\\mu$, $q$ or $C_T$ move it to the empty branch."
              % (len(c1), pm(a1("held_accelerated_efolds")), pm(a1("max_matter_fraction")), pm(a1("lnG")), pm(a1("phi_bounce_over_v")))))
add(("p", "An earlier version of this model prescribed three steps that the present equations produce themselves: the dark-matter abundance was reset at "
          "freeze-out, the hot-phase restoration of the vacuum was applied through its adiabatic result, and the dissipation was capped (hot_cycles.py). That "
          "version gives exactly self-similar cycles with $G_n=%.1f=(S_{n+1}/S_n)^{1/3}$ and %.2f held e-folds, and its controls fail as expected (without the "
          "freeze-out reset no matter era remains; without dissipation the bounce is cold and kinetic dominated). The continuous cycles reproduce its "
          "qualitative structure, and in addition expose the failure that the resets hid (the cold bounce above) and the cycle-to-cycle variation of the held "
          "epoch explained in Section 5.2."
          % (DIL.gm(DIL.GN), DIL.BC[1]["held_accelerated_efolds"])))
rows = []
for c in CR["cycles"][:8] + ([CR["cycles"][-2]] if len(CR["cycles"]) > 9 else []):
    rows.append([str(c["n"]), "%.0f" % (c.get("T_bounce", 0) * MEV) if c.get("T_bounce") else "--", "%.1f" % c["bounce_x"] if c.get("bounce_x") else "--",
                 "%.3f" % c["phi_bounce_over_v"] if c.get("phi_bounce_over_v") is not None else "--", "%.2f" % fo(c) if fo(c) else "(today)",
                 "%.3f" % c["Y_held_over_kappa"] if c.get("Y_held_over_kappa") else "(today)", "%.3f" % c.get("max_matter_fraction", 0),
                 "%.2f" % c.get("held_accelerated_efolds", 0), "%.2f" % c["lnG"] if c.get("lnG") else "--", "%.2f" % c["lnG_entropy"] if c.get("lnG_entropy") else "--"])
add(("table", "Table 2. Continuous cycles without resets (baseline: $m=%s$ MeV, $\\rho_c^{1/4}$ = %.0f MeV, $C_T=1$, $a_c=2$, $\\mu=30H_0$, $q=1$, "
              "$\\Omega_{\\sigma,0}=10^{-30}$). Cycle 0 starts today; $x_b$ and $\\varphi_b$ at the bounce that ends the cycle; $Y/\\kappa$ the dark matter "
              "relative to the present value; the last row is cycle %d. All cycles: continuous_cycles.csv." % (MB, A44["T_bounce_MeV"], CR["cycles"][-2]["n"]),
     ["cycle", "$T_b$ (MeV)", "$x_b$", "$\\varphi_b/v$", "$Y_{\\rm fo}/\\kappa$", "$Y/\\kappa$ held", "peak matter", "held e-folds", "$\\ln G_n$",
      "$\\frac13\\ln\\frac{S_{n+1}}{S_n}$"], rows, None))
add(("fig", "fig17_continuous_cycles.png", "Fig. 7. Continuous cycles without resets. (a) Fractions of radiation, matter and vacuum energy against the bath "
            "temperature in the expanding branch, for today's cycle (dotted, from today) and the next (solid). (b) The field in three cycles: frozen at the "
            "bounce, released from the hot phase, damped to $\\varphi\\approx0$, held by the matter and released. (c) The dark matter of cycle 1 against "
            "$x=m_{\\rm eff}/T$: in equilibrium at the bounce, frozen out while the field is displaced (upper plateau), and reduced to the present ratio when "
            "the restoration lowers its mass and heats the bath (lower plateau). (d) Held e-folds, $Y/\\kappa$ and $\\ln G_n$ over %d cycles." % N_CYC, 6.3))
add(("h2", "5.2 The cycle as an attracting fixed point"))
m10, m100 = CM["map"]["5500000000.0"], CM["map"]["55000000000.0"]
HE = CM.get("held", {})
add(("p", "Why the cycle repeats can be stated as a property of a map. At a bounce the density is fixed at $\\rho_c$, the field is overdamped with its "
          "velocity slaved, and the remaining freedom is the value $\\varphi_b$ at which the field is frozen. The bounce-to-bounce map "
          "$\\varphi_{b,n}\\mapsto\\varphi_{b,n+1}$ was probed by starting one cycle from the baseline bounce state with $\\varphi_b$ set to each of 15 values "
          "between $0.02v$ and $v$ and integrating to the next bounce (cycle_map.py, Fig. 8a). For $m=10$ MeV every start gives a working cycle, the map has a "
          "single fixed point at $\\varphi_b^*=%.3fv$, and its slope there is $%.2f$: a deviation shrinks by a factor of about %.0f in one cycle, so the "
          "repeating cycle is an attracting fixed point whose basin covers the whole interval tested. For $m=100$ MeV the fixed point lies at $%.3fv$ with slope "
          "$%.2f$: still attracting, with oscillating and slower convergence. The empty branch of Section 6.2 is not a second fixed point of this map; it is "
          "reached when the dark matter is not regenerated at the bounce, a second state variable that the one-dimensional map holds fixed."
          % (m10["fixed_point"], m10["slope_at_fixed_point"], 1 / abs(m10["slope_at_fixed_point"]), m100["fixed_point"], m100["slope_at_fixed_point"])))
add(("p", "The held epoch itself has a closed form. Acceleration begins when the matter density falls to $2V_c$ and the field is released when it falls to "
          "$\\rho_{\\rm crit}$, which takes $\\frac13\\ln(2V_c/\\rho_{\\rm crit})$ e-folds; the tachyonic field then needs $N_{\\rm grow}$ more e-folds to grow "
          "from its restored value $\\varphi_{\\rm rest}$ to $0.1v$,"))
add(("eq", r"N_{\rm held}=\frac{1}{3}\ln\frac{2V_c}{\rho_{\rm crit}}+N_{\rm grow},\qquad \int_0^{N_{\rm grow}}\frac{\lambda(N)}{H_\Lambda}\,dN=\ln\frac{0.1\,v}{\varphi_{\rm rest}},\qquad \lambda=\Big(\mu^2\big(1-e^{-3N}\big)+\tfrac94H_\Lambda^2\Big)^{1/2}-\tfrac32H_\Lambda"))
add(("p", "with $H_\\Lambda=(V_c/3)^{1/2}$. With $\\varphi_{\\rm rest}$ read from each cycle of the %d-cycle run, this reproduces the integrated held e-folds to "
          "%s e-folds (rms over %d cycles; largest deviation %s; Fig. 8b). The variation of the held epoch between cycles is therefore explained: it is the "
          "logarithm of how deeply each cycle restored the field, divided by the tachyonic growth rate."
          % (N_CYC, "%.3f" % HE["rms_error"], HE["n_points"], "%.3f" % HE["max_abs_error"])))
add(("fig", "fig18_cycle_map.png", "Fig. 8. (a) Bounce-to-bounce return map of the frozen field value for $m=10$ and 100 MeV (bounce at %.0f MeV); dotted: "
            "the identity, whose intersection with each curve is the repeating cycle. (b) Held e-folds from the integration against the closed form, one point "
            "per cycle." % A44["T_bounce_MeV"], 6.3))

# ============================================================================================ 6 numerics and viability
add(("h1", "6. Numerical validation and parameter viability"))
add(("h2", "6.1 Numerical method and the admissibility of its devices"))
add(("p", "The cycle is integrated with the state $(\\ln a, H, \\varphi, \\dot\\varphi, T, N)$ and two auxiliary integrals (the radiation history and the "
          "pressure-volume work), in units of $H_0$ and $M_{\\rm Pl}$. The integration proceeds in chunks of one doubling or halving of the scale factor with "
          "the clock restarted in each chunk, because near the bounce the time scale $\\rho_c^{-1/2}$ is many orders of magnitude below the cycle time and an "
          "absolute clock cannot resolve it. Absolute tolerances are set relative to each variable, an eighth-order Runge--Kutta method is tried first with "
          "LSODA and Radau as fall-backs, the overdamped reduction is used where $\\Upsilon>10^3\\max(m_{\\rm eff},H,\\mu)$ (%d chunks of the baseline run), and "
          "the constraint is projected away from the turning points (largest correction %s). The Maxwell--Boltzmann integrals use the asymptotic series of the "
          "Bessel functions for $x>100$, where the library functions lose precision, and the heat capacity of the gas is computed analytically."
          % (CR.get("overdamped_chunks", 0), sci(CR["projections"]["max_rel"], 0))))
others = [r for r in NA["rows"] if r["variant"] != "baseline"]
worst = lambda r, k: r["max_rel_change_vs_baseline"].get(k, float("nan"))
okv = [r for r in others if r["classification"] == "working"]
bad = [r for r in others if r["classification"] != "working"]
res_of = {r["variant"]: r["max_residual"] for r in NA["rows"]}
scat = NA["baseline_cycle_scatter"]
add(("p", "The results do not depend on how the numerical devices are set (numerics_admissibility.py, three cycles per variant). A tighter tolerance "
          "($10^{-10}$), a ten times stricter threshold for the overdamped reduction, and the constraint projection applied only at every fourth chunk "
          "change the held e-folds by at most %.2g, the held-epoch abundance by at most %.2g and $\\ln G_n$ by at most %.2g (relative), against a scatter "
          "between cycles of %.2g, %.2g and %.2g in the baseline itself; every variant is a working cycle. The projection cannot be removed altogether: "
          "without it the Friedmann-constraint residual grows to %s through the stiff bounce and contraction phases and the integration stops before the "
          "first cycle closes. Projection keeps the integrator on the constraint surface, which the exact equations never leave, while its frequency does not "
          "matter (the residual rises from %s to %s when it is applied four times less often)."
          % (max(worst(r, "held_accelerated_efolds") for r in okv), max(worst(r, "Y_held_over_kappa") for r in okv), max(worst(r, "lnG") for r in okv),
             scat["held_accelerated_efolds"], scat["Y_held_over_kappa"], scat["lnG"], sci(bad[0]["max_residual"], 0) if bad else "--",
             sci(res_of["baseline"], 0), sci(res_of.get("projection every 4th chunk", 1.0), 0))))
ISR = {r["variant"]: r for r in ISV["rows"]}
ISok = [r for r in ISV["rows"] if r["settings"].get("gravity") == "constrained" and r["status"] == "completed"]
ISall = [r for r in ISV["rows"] if r["settings"].get("gravity") == "constrained"]
iw = lambda k: max(r["max_rel_change_vs_projected"].get(k, 0.0) for r in ISok)
add(("p", "Two diagnostics must not be confused. The residual is the violation of the Friedmann constraint measured inside each chunk, normalized by the "
          "sum of the absolute values of the energy components; the projection correction is the relative change of $H$ applied at the start of a chunk. "
          "They differ by normalization, since $H$ is small wherever the components nearly cancel, and a small residual after projection does not by itself "
          "show that the unprojected trajectory is accurate. Table 3 lists both for every setting, together with the changes of the turnaround (its scale "
          "factor and the duration of the cycle) and of the cycle properties."))
add(("p", "The decisive test is an independent formulation of the gravitational equations that needs no projection (independent_solver.py). Away from the "
          "turning points $H$ is not integrated but computed from the constraint, $H=\\pm[\\rho(1-\\rho/\\rho_c)/3]^{1/2}$, so the constraint holds identically; "
          "close to a turnaround or bounce, where the square root carries no information about the change of sign, $H$ is evolved across a short bridge with "
          "the Raychaudhuri equation, entered when $H^2$ falls below a fixed fraction of the natural rate, and the constraint mismatch accumulated across each "
          "bridge is recorded. With the same matter equations, %d of %d settings of this solver complete three cycles; their bridges accumulate a mismatch of "
          "at most %s, and they reproduce the projected integration to %s in the turnaround scale factor, %s in the duration of the cycle, %s in the held "
          "e-folds, %s in the held-epoch abundance and %s in $\\ln G_n$. The projection is therefore not what makes the cycle repeat."
          % (len(ISok), len(ISall), sci(max(r["projections"]["max_rel"] for r in ISok) + 1e-16, 0), sci(iw("ln_a_max"), 0), sci(iw("t_cycle"), 0),
             sci(iw("held_accelerated_efolds"), 0), sci(iw("Y_held_over_kappa"), 0), sci(iw("lnG"), 0))))
SHORTV = {"baseline": "baseline", "rtol 1e-10": "rtol $10^{-10}$", "overdamped threshold 1e4": "overdamped $10^4$",
          "projection every 4th chunk": "projection every 4th", "no projection": "no projection"}
nrows = []
for r in NA["rows"]:
    d = r.get("max_rel_change_vs_baseline", {})
    g = lambda k: (sci(d[k], 0) if d.get(k) else ("0" if k in d else "--"))
    nrows.append(["P: " + SHORTV.get(r["variant"], r["variant"]), "failed" if r["classification"] != "working" else "working",
                  sci(r["max_residual"], 0), "%d" % r["projections"]["n"], sci(r["projections"]["max_rel"], 0) if r["projections"]["max_rel"] else "--",
                  g("ln_a_max"), g("t_cycle"), g("held_accelerated_efolds"), g("lnG")])
for r in ISV["rows"]:
    if r["settings"].get("gravity") != "constrained":
        continue
    d = r["max_rel_change_vs_projected"]
    g = lambda k: (sci(d[k], 0) if d.get(k) else "--")
    nrows.append(["C: " + r["variant"].split(", ", 1)[1].replace("rtol 1e-10", "rtol $10^{-10}$"), "working" if r["status"] == "completed" else "stopped",
                  sci(r["max_residual"], 0), "%d" % r["projections"]["n"], sci(r["projections"]["max_rel"] + 1e-16, 0),
                  g("ln_a_max"), g("t_cycle"), g("held_accelerated_efolds"), g("lnG")])
add(("table", "Table 3. Numerical diagnostics of the baseline (three cycles per setting; P: projected solver, C: constrained solver with the bridge "
              "parameter given). Residual: largest Friedmann-constraint violation inside a chunk; "
              "corrections: number and largest relative size of the projections of $H$ (projected solver) or the constraint mismatch at the end of the "
              "Raychaudhuri bridges (constrained solver); the last four columns give the largest relative change from the projected baseline in the "
              "turnaround scale factor, the duration of the cycle, the held e-folds and $\\ln G_n$. The no-projection run stopped before the first cycle closed.",
     ["Setting", "Outcome", "Residual", "Count", "Max. corr.", "$\\ln a_{\\max}$", "$t_{\\rm cycle}$", "Held", "$\\ln G_n$"], nrows, None))
add(("h2", "6.2 Where the cycles work, and how they fail"))
grid, marg, robust = CYB.grid, CYB.marg, CYB.robust
short, lab, lab_r = CYB.short, CYB.lab, CYB.lab_r
n_rob = sum(short(r) == "working" for r in robust)
add(("p", "Over a grid of dark-matter mass (1 MeV to 1 GeV) and bounce temperature (1.8 MeV to 1.8 GeV), with five cycles per run, the cycles work in %d of "
          "%d cells (Fig. 6b). Every failure in the grid is F1: the bounce is too cold to return the dark matter to equilibrium, the dark matter is annihilated in "
          "the contraction, and the next cycle has no matter to hold its vacuum. The dynamics are bistable: the field freezes at the bounce either at "
          "$\\varphi_b\\approx0.23$--$0.31v$, where the dark matter is light enough to rethermalize and the cycle reproduces itself, or near $v$, where it is "
          "$A(v)\\approx85$ times heavier and is not regenerated. The condition that does not depend on the phase of the frozen field is a bounce that "
          "rethermalizes even the heaviest dark matter, $T_b\\gtrsim mA(v)/5$. For $m=100$ MeV and $T_b=180$ MeV it is not met, and varying one parameter at a "
          "time about that point gives: %s. For $m=10$ MeV at the same bounce it is met, and %d of %d variations work: %s. The classification looks for F1 (dark "
          "matter not regenerated), F2 (vacuum not restored), F3 (bounce dominated by the scalar's kinetic energy), F4 (bounce dominated by shear), F5 (no "
          "bounce) and F6 (numerical)."
          % (CYB.n_work, len(grid), "; ".join("%s: %s" % (lab(r["params"]), short(r)) for r in marg if "bounce" not in r["params"]), n_rob, len(robust),
             "; ".join("%s: %s" % (lab_r(r), short(r)) for r in robust))))
add(("h2", "6.3 Joint viability over parameter ranges"))
add(("p", "How often are all the conditions met at once? Over stated ranges (viability_mc.py; %d draws, log-uniform $a_c\\in[0.5,10]$, $m\\in[1,1000]$ MeV, "
          "$T_b\\in[1\\,{\\rm MeV},10\\,{\\rm GeV}]$, uniform $q\\in[-1,3]$) the vacuum is held today in %.0f%% of draws, recollapses in %.0f%%, has its bounce in "
          "the window between BBN and black-hole percolation in %.0f%%, and rethermalizes the dark matter, $T_b>mA(v)/5$ with the exact $A(v)$ of Section 2.1, "
          "in %.0f%%. All four hold together in %.1f%% of the draws, %.1f%% of those that are held today and recollapse; rethermalization is the binding "
          "condition. The window uses the baseline estimate of the black-hole bound; with the lower estimates of Section 4.5 the joint fraction falls to "
          "%.1f--%.1f%%. Dissipation, a coherent seed and inherited rather than regenerated shear are required in addition and are not sampled. The fractions "
          "describe the sample space of the model, not a probability for our universe."
          % (HM["n"], 100 * HM["single"]["H1 held today"], 100 * HM["single"]["H2 recollapse"], 100 * HM["single"]["H3 bounce window"],
             100 * HM["single"]["H4 re-thermalization"], 100 * HM["joint"], 100 * HM["joint_given_H1H2"],
             100 * RS["held_viability"]["region overdense by 100, overlap"]["joint"], 100 * RS["held_viability"]["overlap threshold"]["joint"])))
add(("h2", "6.4 Executable checks"))
DESC = {"V27": "matter-held vacuum: held today, release, recollapse; control fails", "V28": "prescribed-step hot cycles (earlier version) and their controls",
        "V29": "continuous cycles: 100 cycles with dark matter regenerated, matter era, held epoch; growth = entropy$^{1/3}$; cold bounce fails",
        "V30": "causal test: frozen matter never recollapses; diluting model recollapses for every $q>0$",
        "V31": "shear scaling and ekpyrotic gain; bounce invariants and continuity; black-hole volume fraction",
        "V32": "attracting fixed point; closed-form held epoch; numerical devices; independent solver", "V33": "phase structure (spinodal release) and coherence of the release"}
vrows = []
for k in P1_CHECKS:
    v = V[k]
    det = {kk: (("%.3g" % vv) if isinstance(vv, float) else str(vv)) for kk, vv in v.items() if kk != "passed"}
    vrows.append([k.split("_")[0], DESC.get(k.split("_")[0], k), "pass" if v["passed"] else "FAIL"])
add(("p", "verify.py recomputes the checks of Table 4 from the stored results; the manuscript is built only if they pass. The repository also contains "
          "checks of an earlier tunneling analysis and of the companion interface paper, which are not used here."))
add(("table", "Table 4. Executable checks of this paper (verify.py).", ["No.", "What is checked", "Outcome"], vrows, None))

# ============================================================================================ 7 observations and limits
add(("h1", "7. Observational constraints and limitations"))
add(("h2", "7.1 Observational consequences"))
MODES, DS = CYB.MODES, CYB.DS
m0 = MODES[0]
add(("bullets", [
    "Dark energy. While held, the vacuum has $1+w<%s$; the model predicts no evolution of the dark energy before the release, the release near $a=a_c$ and "
    "a turnaround %.0f Gyr from now for $a_c=2$, $\\mu=30H_0$, $q=1$. A confirmed evolution of $w$ at the level suggested by recent baryon-acoustic-"
    "oscillation analyses \\cite{DESI} would exclude it." % (texsci(DS["one_plus_w_today_bound"], 0), DS["t_turn_Gyr"]),
    "Fifth forces and voids. The coupling vanishes at $\\varphi=0$, so the homogeneous background and dense, collapsed structures carry no fifth force, "
    "and dark-matter perturbations do not source the field at linear order (Section 2.1). This is not a full equivalence with $\\Lambda$CDM: the growth "
    "of structure at second order and in the time-dependent mass $m_{\\rm eff}$ has not been computed. A void with density below $\\rho_{\\rm dm,0}/a_c^3$ "
    "and radius above the Compton length $\\pi/\\mu$ would release the vacuum locally today; void statistics bound the parameter space for large $\\mu$ and "
    "small $a_c$.",
    "Dark matter. The cycle requires a thermal relic rethermalized at the bounce: $m\\lesssim5T_b/A(\\varphi_b)$ and, with Section 4.5, $m$ below a few "
    "hundred MeV for $T_b\\lesssim%.1f$ GeV; a nonthermal dark matter (an axion produced by misalignment, for instance) or a heavier relic is incompatible "
    "with the cycle; thermal relics below a few MeV are separately constrained by BBN and the CMB \\cite{Sabti20}." % (RL["T_perc_E3_MeV"] / 1e3),
    "Primordial perturbations. The scales observed today, $k=(1$--$1000)a_0H_0$, cross the Hubble radius of the preceding contraction at $T=%.0f$--$%.0f$ "
    "(in units of 1.82 meV), where the contraction is radiation dominated ($w=%.3f$). A contraction with constant $w$ gives $n_s-1=12w/(1+3w)$ "
    "\\cite{Wands99,FB02}, i.e. $n_s=%.2f$, with an amplitude set by $\\rho_c/M_{\\rm Pl}^4=%s$, against $n_s=0.965$ and $2.1\\times10^{-9}$ observed "
    "\\cite{PlanckX}. The background cycle does not produce the observed perturbations; a separate mechanism (for example a curvaton, or an entropic "
    "mechanism in an ekpyrotic phase) would be required, and none is part of this model."
    % (m0["T_exit"], MODES[-1]["T_exit"], m0["w_exit"], m0["n_s"], texsci(PR["perturbations"]["rho_c_over_MPl4"], 0)),
    "Gravitational waves. Tensor modes from the contraction have the same tiny amplitude; the release is a homogeneous roll without bubble collisions; no "
    "primordial gravitational-wave signal is predicted."]))
add(("h2", "7.2 Consistency conditions and falsification tests"))
add(("p", "The conditions the model must satisfy are of two kinds. Internal conditions, if violated, make the model ill defined; observational tests, if "
          "failed, disfavor it."))
add(("bullets", [
    "C1 (held vacuum). The vacuum is held today only if $a_c>1$, i.e. $\\mu^2M^2<\\rho_{\\rm dm,0}$, and a release into a negative vacuum needs "
    "$A(v)-1=2(V_c-V_0(v))/\\rho_{\\rm crit}>2V_c/\\rho_{\\rm crit}$, a dark-matter mass that grows severalfold after the release (Section 2.1).",
    "C2 (coherent release). The roll must be seeded coherently over the region, by a tilt $\\epsilon\\gtrsim%s\\,\\mu^2v$ or an equivalent seed; otherwise "
    "domain walls form (Section 2.2)." % texsci(COH["eps_hat_min"], 0),
    "C3 (dissipation). Repeated cycles are stable only if the dissipation per cycle exceeds the logarithm of the amplification of the scalar energy, "
    "$\\Upsilon P\\gtrsim\\ln\\mu$ (Section 4.2).",
    "C4 (bounce window). The bounce must lie between BBN (5 MeV) and the percolation of the black holes of the previous cycle, estimated at %.1f--%.1f GeV "
    "within the population prescription and lower if accretion is much faster than geometric capture or the contraction strongly inhomogeneous (Section 4.5)."
    % tuple(RS["physical_accretion_range_GeV"]),
    "C5 (rethermalization). The bounce must return the dark matter to equilibrium, $T_b\\gtrsim mA(v)/5$; otherwise the next cycle is empty (Sections 5.1, 6.2).",
    "C6 (anisotropy). Shear regenerated in a cycle must stay below $%s$ of the matter density at the turnaround, or an ekpyrotic phase must dilute it "
    "(Section 4.3)." % texsci(NOEK["s_max"], 0),
    "O1 (dark energy). $w=-1$ exactly until the release; a confirmed evolution of $w$ excludes the model.",
    "O2 (dark matter). A thermal relic below a few hundred MeV; a nonthermal or heavier dark matter excludes the cycle.",
    "O3 (voids). The emptiest large voids must not release the vacuum today, which bounds large $\\mu$ at small $a_c$.",
    "O4 (perturbations). The observed spectrum must come from a mechanism outside the model; the model makes no prediction for $n_s$, $f_{\\rm NL}$ or $r$."]))
add(("h2", "7.3 What is shown, and what is assumed"))
add(("table", "Table 5. Status of the cyclic model.", ["Element", "Status", "Where"],
     [["Vacuum held by matter, released by dilution", "derived within the model; intervention test passed", "2, 3"],
      ["Turnaround", "follows from $V_0(v)<0$; a negative vacuum alone also recollapses", "3.1, 3.3"],
      ["Coherent release", "needs a seed or tilt $\\epsilon\\gtrsim%s\\,\\mu^2v$" % texsci(COH["eps_hat_min"], 0), "2.2"],
      ["Bounce", "assumed (effective, far below the Planck density); a ghost-fluid bounce gives cycles but not self-similar ones", "4.1, 4.4, 4.6"],
      ["Reheating", "continuous, no reset: dissipation $\\Upsilon=C_TT^3$ (assumed form) and Boltzmann dark matter", "5.1"],
      ["Dark-matter abundance", "output of freeze-out in every cycle; $\\langle\\sigma v\\rangle$ fixed once by today", "5.1"],
      ["Vacuum restoration", "output (frozen at the bounce, damped in the expansion)", "5.1, 5.2"],
      ["Entropy", "grows each cycle; absorbed by Tolman growth $G_n=(S_{n+1}/S_n)^{1/3}$", "5.1"],
      ["Anisotropy", "inherited shear diluted by $G_n^6$; regenerated shear needs $s<10^{-27}$ or ekpyrosis", "4.3"],
      ["Black holes", "estimated bound $T_b\\lesssim$ %.1f--%.1f GeV (prescription dependent); relics diluted" % tuple(RS["physical_accretion_range_GeV"]),
       "4.5"],
      ["$\\Lambda$CDM past", "homogeneous background only; perturbations and voids not tested", "2.1, 7.1"],
      ["Inhomogeneities, perturbations", "not followed; the observed spectrum is not produced", "7.1"],
      ["Joint viability", "%.1f%% of the sampled parameter space" % (100 * HM["joint"]), "6.3"]], None))
add(("p", "The limits are definite. The bounce is an effective prescription, and $\\rho_c$ lies far below the Planck density in the simulations; the "
          "dissipation has an assumed form; the dynamics are homogeneous, so the growth of inhomogeneities in the contraction, which can dominate a contracting "
          "universe, is not followed (Bianchi I shear is the only anisotropy treated); the dark matter is a single species with Maxwell--Boltzmann statistics; "
          "and $a_c\\sim1$ requires $\\mu^2M^2\\approx\\rho_{\\rm dm,0}$, a coincidence of the same kind as that of the cosmological constant. A first-order "
          "version of the trigger, in which the held vacuum becomes metastable and decays by bubble nucleation, is a distinct extension with its own questions "
          "and is not needed for the cycle reported here."))

# ============================================================================================ 8 conclusions
add(("h1", "8. Conclusions"))
add(("p", "Within a stated model, the dilution of matter can end the cosmological expansion and, with an effective bounce, start a new hot phase. A scalar "
          "coupled to dark matter holds the observed dark energy with a homogeneous $\\Lambda$CDM background and releases it by dilution into a negative vacuum; the expansion turns "
          "around, and an intervention that freezes the density seen by the field removes the turnaround. With an effective loop-quantum-cosmology bounce, "
          "dissipation and Boltzmann dark matter, %d consecutive cycles were followed with no variable reset; in each, the dark matter is rethermalized and freezes "
          "out to the observed ratio, the vacuum is restored, held and released, and the entropy produced enlarges the next cycle by $(S_{n+1}/S_n)^{1/3}$. The "
          "repeating cycle is an attracting fixed point of the bounce-to-bounce map, and its dark-energy epoch has a closed form."
          % N_CYC))
add(("p", "The conditions are explicit: a bounce at a fixed density, hot enough to rethermalize the dark matter and cold enough to keep the black holes of "
          "the previous cycle from percolating; dissipation; a thermal-relic dark matter; an anisotropy that is inherited rather than regenerated; and a coherent "
          "seed for the release. Over stated parameter ranges they hold together in %.1f%% of the sample space. With a colder bounce the cycle fails. The model "
          "does not produce the observed primordial perturbations, so these are conditional background solutions: they show what a cyclic hot big bang "
          "triggered by dilution requires, not that the universe has one." % (100 * HM["joint"])))
add(("p", "Three extensions would test the framework further: a derivation of the bounce, or its replacement by a mechanism that fixes the bounce density "
          "from first principles; an inhomogeneous simulation of the contraction, in which the growth of structure rather than homogeneous shear decides "
          "whether the bounce is smooth; and a source of the primordial perturbations compatible with the cycle, for example an entropic mechanism in a "
          "contracting ekpyrotic phase. Each can falsify the model as it stands."))

# ============================================================================================ declarations and references
add(("h1", "Declarations"))
add(("pb", "Funding.", "This research received no specific grant from funding agencies in the public, commercial, or not-for-profit sectors."))
add(("pb", "Conflict of interest.", "The author declares no competing interests."))
add(("pb", "Data availability.", "The data underlying every figure and table are in the repository listed below, as JSON and CSV files written by the scripts "
        "named in the text."))
add(("pb", "Code availability.", "All code (Python, NumPy, SciPy, Matplotlib \\cite{NumPy,SciPy}) is at %s (MIT license) and archived at https://doi.org/%s; "
        "this manuscript is archived at https://doi.org/%s. Running the scripts in the order given in the README regenerates every number, table and figure."
        % (GITHUB, DOI_CODE, DOI_PAPER)))
add(("pb", "Related manuscript.", "A companion manuscript on density-dependent vacuum bubble nucleation in the same class of matter-coupled models is "
        "in preparation for separate submission to Physics of the Dark Universe%s; it shares the class of models but none of the results reported here."
        % ((", archived at https://doi.org/%s" % COMPANION_DOI) if COMPANION_DOI else "")))
add(("pb", "Use of generative AI.", "During the preparation of this work, the author used Claude (Anthropic) to assist with the design of the numerical "
        "analysis, code, literature checking and drafting. The author reviewed and edited all content and takes full responsibility for it."))
if COMPANION_DOI:
    REFS["Companion"] = ("L. Sandler, Density-dependent vacuum bubble nucleation in matter-coupled dark-sector cosmology: wall dynamics, transmission and "
                         "gravitational viability. Preprint, Zenodo (2026). https://doi.org/%s" % COMPANION_DOI)
add(("h1", "References"))

# number citations in order of first appearance
order = []
for blk in B:
    chunks = []
    if blk[0] in ("p", "abstract", "pb"):
        chunks = [blk[-1]]
    elif blk[0] == "bullets":
        chunks = list(blk[1])
    elif blk[0] == "table":
        chunks = [str(c) for row in blk[3] for c in row]
    for ch in chunks:
        for m in re.finditer(r"\\cite\{([^}]+)\}", ch):
            for key in m.group(1).split(","):
                if key not in order:
                    order.append(key)
undefined = [k for k in order if k not in REFS]
if undefined:
    sys.exit("undefined citation keys %s" % undefined)
numbering = {k: i + 1 for i, k in enumerate(order)}
rep = lambda s: re.sub(r"\\cite\{([^}]+)\}", lambda m: "[" + ", ".join(str(numbering[x]) for x in m.group(1).split(",")) + "]", s)
for k in order:
    add(("ref", "[%d] %s" % (numbering[k], REFS[k])))
B2 = []
for b in B:
    if b[0] in ("p", "abstract", "note"):
        B2.append((b[0], rep(b[1])))
    elif b[0] == "bullets":
        B2.append(("bullets", [rep(x) for x in b[1]]))
    elif b[0] == "pb":
        B2.append(("pb", b[1], rep(b[2])))
    elif b[0] == "table":
        B2.append(("table", b[1], b[2], [[rep(str(c)) for c in row] for row in b[3]], b[4]))
    else:
        B2.append(b)
B = [CE.block(b) for b in B2]
# figures are numbered by order of appearance: map the numbers written in the captions to their order, then rewrite every "Fig./Figs. n" reference
_fig_old = [int(re.match(r"Fig\. (\d+)\.", b[2]).group(1)) for b in B if b[0] == "fig"]
FIGMAP = {old: i + 1 for i, old in enumerate(_fig_old)}
assert len(FIGMAP) == len(_fig_old), "duplicate figure numbers in captions"


def _refig(t):
    return re.sub(r"(Figs?\. )(\d+)", lambda m: m.group(1) + str(FIGMAP.get(int(m.group(2)), int(m.group(2)))), t) if isinstance(t, str) else t


def _refig_block(b):
    k = b[0]
    if k in ("p", "abstract", "note"):
        return (k, _refig(b[1]))
    if k == "bullets":
        return (k, [_refig(x) for x in b[1]])
    if k == "pb":
        return (k, b[1], _refig(b[2]))
    if k == "table":
        return (k, _refig(b[1]), b[2], [[_refig(str(c)) for c in row] for row in b[3]], b[4])
    if k == "fig":
        return (k, b[1], _refig(b[2])) + tuple(b[3:])
    return b


B = [_refig_block(b) for b in B]


def fix_sci(text):
    if "⟦" not in text:
        return text
    parts, out = text.split("$"), []
    for i, seg in enumerate(parts):
        out.append(seg.replace("⟦", "").replace("⟧", "") if i % 2 == 1 else seg.replace("⟦", "$").replace("⟧", "$"))
    return "$".join(out)


def map_strings(block):
    k = block[0]
    if k in ("p", "abstract", "note", "title"):
        return (k, fix_sci(block[1]))
    if k == "pb":
        return (k, block[1], fix_sci(block[2]))
    if k == "bullets":
        return (k, [fix_sci(x) for x in block[1]])
    if k == "table":
        return (k, fix_sci(block[1]), [fix_sci(h) for h in block[2]], [[fix_sci(str(c)) for c in row] for row in block[3]], block[4])
    if k == "fig":
        return (k, block[1], fix_sci(block[2])) + tuple(block[3:])
    return block


B = [map_strings(b) for b in B]
BASENAME = "EPJC_Dilution_Triggered_Cycles"


def build_letter():
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11); st.paragraph_format.space_after = Pt(8)
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(1.1); s.top_margin = s.bottom_margin = Inches(1.0)

    def para(text, bold=False, after=None):
        p = d.add_paragraph(); r = p.add_run(CE.ce(text)); r.bold = bold
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
    para("Leon Sandler", True, 0)
    for t in ("Independent Researcher", "Northbrook, Illinois, USA", "sandler.leon@gmail.com", "ORCID 0009-0007-4584-808X"):
        para(t, after=0)
    para("")
    para(datetime.date.today().strftime("%d %B %Y").lstrip("0"))
    para("The Editors", after=0); para("The European Physical Journal C", after=0); para("Springer Nature")
    para("Dear Editors,")
    para("I submit the manuscript “%s” as a Regular Article (Theoretical Physics) for consideration in The European Physical Journal C." % TITLE)
    para("In plain terms, the paper reverses a common intuition about the far future. The dilution of matter is usually the end state of cosmic evolution, "
         "the approach to heat death; here it is the trigger. A scalar field coupled to dark matter holds a vacuum with the observed dark-energy density while "
         "the dark-matter density is high, with a homogeneous background that reproduces LambdaCDM, and releases it by dilution into a negative vacuum, which "
         "ends the expansion. An intervention that freezes the matter density seen by the field removes the turnaround, and the uncoupled field with the same "
         "potential cannot reproduce the observed past over most of the parameter range.")
    para("With an effective loop-quantum-cosmology bounce, thermal dissipation and dark matter that obeys its own Boltzmann equation with a field-dependent "
         "mass, %d consecutive cycles are integrated without resetting any variable. Each has a radiation-dominated bounce, freeze-out to the observed "
         "abundance, a matter era, a dark-energy epoch, a release and a turnaround, and the repeating cycle is an attracting fixed point of a bounce-to-bounce "
         "map. The paper states what the cycle requires and where it fails: a bounce hot enough to rethermalize the dark matter but below the percolation of "
         "the previous cycle's black holes, dissipation, inherited rather than regenerated anisotropy and a coherent seed for the release; over stated parameter "
         "ranges these hold together in %.1f%% of the sample space." % (N_CYC, 100 * HM["joint"]))
    para("The claims are limited deliberately. The bounce is an effective prescription, the dynamics are homogeneous, and the cycle does not produce the "
         "observed primordial perturbations (it gives n_s near 3), so the results are conditional background solutions, not an observationally complete "
         "cosmology. The ingredients (symmetron-type couplings, coupled dark matter, effective bounces, warm-inflation dissipation, freeze-out, Tolman's "
         "growing cycles) are established and cited; what is new is their combination into a cycle followed without resets, the exact relation that fixes the "
         "growth of the dark-matter mass at the release, the intervention tests, and the map of failure conditions.")
    para("Related submission. An earlier, broader version of this work also contained an analysis of localized vacuum bubbles. Following a review, that part "
         "has been separated into an independent companion manuscript on density-dependent vacuum bubble nucleation, which is in preparation for separate "
         "submission to Physics of the Dark Universe%s. The two papers share the class of models; neither relies on the results of the other, and I will "
         "provide the companion manuscript on request." % ((" (preprint https://doi.org/%s)" % COMPANION_DOI) if COMPANION_DOI else ""))
    para("The analysis is fully reproducible. The code (Python, MIT license) and the data behind every figure and table are at %s and archived at "
         "https://doi.org/%s; the manuscript is archived at https://doi.org/%s. The manuscript is built from the stored results and only if the executable "
         "checks pass." % (GITHUB, DOI_CODE, DOI_PAPER))
    para("A generative AI tool (Claude, Anthropic) was used to assist with the numerical analysis, code, literature checking and drafting; I reviewed all of "
         "it and take full responsibility for the content, as disclosed in the manuscript. The manuscript is original, has not been published elsewhere and is "
         "not under consideration at any other journal. I am the sole author, I have no competing interests, and the work received no external funding.")
    para("Thank you for considering this submission.")
    para("Sincerely,")
    para("Leon Sandler", True)
    d.save(os.path.join(OUT, "EPJC_Cover_Letter.docx"))


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    RN.render_docx(B, os.path.join(OUT, BASENAME + ".docx"), FIGDIR)
    RN.render_tex(B, os.path.join(OUT, BASENAME + ".tex"))
    build_letter()
    words = sum(len(b[1].split()) for b in B if b[0] in ("p", "abstract")) + sum(len(x.split()) for b in B if b[0] == "bullets" for x in b[1])
    nfig = sum(1 for b in B if b[0] == "fig")
    print("Paper 1: %d words, %d figures, %d tables, %d equations, %d references" % (words, nfig, sum(1 for b in B if b[0] == "table"),
                                                                                    sum(1 for b in B if b[0] == "eq"), len(order)))
